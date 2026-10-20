"""Edições conferidas para os gen_*.py do deck (extraído do gen_pendencias de 2026-10-16).

Toda função confere o estado antigo antes de mexer (assert) e registra em ed.log.

    sys.path.insert(0, r"C:\\projetos\\pll_stability_9bus\\.claude\\skills\\tcc-pptx-editor\\scripts")
    from edicoes import Edicao

    ed = Edicao(r"C:\\Temp\\defesa_<tema>")          # lê deck.pptx

    def g_txt(ed):
        ed.troca_run(10, "Text 22", "28 cenários", "29 cenários")

    ed.rodar(globals())      # CLI: gen_x.py [txt,img] [saída.pptx]; padrão: todos os g_*, montado.pptx

Posicionar e escrever formas (subscrito _{x}, caixa, cópia de grupo): montagem.py.

Grupos (g_<nome>) existem para a bisseção: arquivo recusado pelo PowerPoint →
rodar um grupo por vez numa saída própria e renderizar cada variante.
"""
import copy
import sys

from PIL import Image

from pptx_util import EMU, abrir, formas, utf8

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class Edicao:
    def __init__(self, pasta, entrada="deck.pptx"):
        utf8()
        self.pasta = pasta
        self.prs = abrir(f"{pasta}\\{entrada}")
        self.S = list(self.prs.slides)
        self.log = []

    # localização ---------------------------------------------------------
    def forma(self, n, nome, contem=None):
        """Única forma do slide n (base 1) com esse nome e, se dado, esse trecho de texto."""
        c = [sh for sh, _ in formas(self.S[n - 1]) if sh.name == nome and
             (contem is None or contem in (sh.text_frame.text if sh.has_text_frame else ""))]
        assert len(c) == 1, f"slide {n} '{nome}' ({contem}): {len(c)} formas"
        return c[0]

    def tabela(self, n):
        t = [sh for sh, _ in formas(self.S[n - 1]) if getattr(sh, "has_table", False) and sh.has_table]
        assert len(t) == 1, f"slide {n}: {len(t)} tabelas"
        return t[0]

    # texto -----------------------------------------------------------------
    def troca_run(self, n, nome, velho, novo):
        """Troca dentro de um run, preservando o estilo. O texto antigo tem de estar em exatamente 1 run."""
        sh = self.forma(n, nome, velho)
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs if velho in r.text]
        assert len(runs) == 1, f"slide {n} '{nome}': '{velho}' em {len(runs)} runs"
        runs[0].text = runs[0].text.replace(velho, novo)
        self.log.append(f"slide {n} {nome}: '{velho}' -> '{novo}'")

    @staticmethod
    def tira_run(run):
        """Remove um run e os a:pPr que sobrarem além do primeiro (armadilhas.md: o
        PowerPoint recusa o arquivo com dois a:pPr seguidos no parágrafo)."""
        p = run._r.getparent()
        p.remove(run._r)
        for extra in p.findall(A + "pPr")[1:]:
            p.remove(extra)

    def notas(self, n, texto, contem):
        """Reescreve as notas do slide n; 'contem' prova que é o slide certo."""
        tf = self.S[n - 1].notes_slide.notes_text_frame
        assert contem in tf.text, f"slide {n}: notas não contêm '{contem}'"
        tf.text = texto
        self.log.append(f"slide {n} notas: {len(texto.split())} palavras")

    # imagem ----------------------------------------------------------------
    def troca_img(self, n, nome, png):
        """Mídia nova na mesma caixa, encaixada pela proporção e centrada."""
        sh = self.forma(n, nome)
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
        _, rid = self.S[n - 1].part.get_or_add_image_part(png)
        sh._element.blipFill.blip.rEmbed = rid
        pw, ph = Image.open(png).size
        nw, nh = (w, int(w * ph / pw)) if pw / ph > w / h else (int(h * pw / ph), h)
        sh.left, sh.top, sh.width, sh.height = x + (w - nw) // 2, y + (h - nh) // 2, nw, nh
        self.log.append(f"slide {n} {nome}: imagem <- {png.split(chr(92))[-1]} "
                        f"({nw / EMU:.2f}x{nh / EMU:.2f} in, caixa {w / EMU:.2f}x{h / EMU:.2f})")

    # slide -----------------------------------------------------------------
    def duplica_slide(self, n):
        """Cópia do slide n (formas, mídia e árvore das notas) anexada no fim; devolve o número novo.
        Mover para a posição certa com move_slides. Ver armadilhas.md (notas sem corpo)."""
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        src = self.S[n - 1]
        novo = self.prs.slides.add_slide(src.slide_layout)
        for ph in list(novo.placeholders):
            ph._element.getparent().remove(ph._element)
        for sh in src.shapes:
            el = copy.deepcopy(sh._element)
            novo.shapes._spTree.append(el)
            if sh.shape_type == 13:  # imagem: religar a mídia no slide novo
                rid = novo.part.relate_to(src.part.related_part(sh._element.blipFill.blip.rEmbed), RT.IMAGE)
                el.blipFill.blip.rEmbed = rid
        if src.has_notes_slide:  # o notes_slide novo nasce sem placeholder de corpo
            arv = novo.notes_slide.shapes._spTree
            for el in list(arv)[2:]:
                arv.remove(el)
            for el in list(src.notes_slide.shapes._spTree)[2:]:
                arv.append(copy.deepcopy(el))
        self.S.append(novo)
        self.log.append(f"slide {len(self.S)}: cópia do slide {n}")
        return len(self.S)

    def renumera_ids(self, n):
        """Ids de forma repetidos no slide n (cópia de forma por deepcopy) viram ids novos."""
        vistos, prox = set(), 1000
        for el in self.S[n - 1].shapes._spTree.iter():
            if el.tag.endswith("}cNvPr"):
                if el.get("id") in vistos:
                    el.set("id", str(prox))
                    prox += 1
                vistos.add(el.get("id"))
        self.log.append(f"slide {n}: {prox - 1000} ids de forma renumerados")

    def move_slides(self, ordem):
        """Nova ordem pelos números atuais (base 1), ex. [1, 2, 4, 3]; precisa listar todos."""
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        assert sorted(ordem) == list(range(1, len(ids) + 1)), "ordem precisa listar todos os slides"
        for el in ids:
            lst.remove(el)
        for k in ordem:
            lst.append(ids[k - 1])
        self.log.append(f"ordem: {ordem}")

    # tabela ----------------------------------------------------------------
    def linha_nova(self, n, modelo, depois, textos):
        """Copia a linha 'modelo' (base 0, cabeçalho = 0) para depois da linha 'depois'."""
        gf = self.tabela(n)
        trs = gf.table._tbl.tr_lst
        trs[depois].addnext(copy.deepcopy(trs[modelo]))
        linha = gf.table.rows[depois + 1]
        assert len(linha.cells) == len(textos), f"slide {n}: {len(linha.cells)} colunas"
        for c, t in zip(linha.cells, textos):
            runs = c.text_frame.paragraphs[0].runs
            runs[0].text = t
            for r in runs[1:]:
                r.text = ""
        gf.height = sum(r.height for r in gf.table.rows)
        self.log.append(f"slide {n} tabela: linha nova após {depois}: {textos}")

    def alturas(self, n, cab, dado):
        """Altura das linhas em polegadas; a forma não cresce sozinha no arquivo."""
        gf = self.tabela(n)
        for i, r in enumerate(gf.table.rows):
            r.height = int((cab if i == 0 else dado) * EMU)
        gf.height = sum(r.height for r in gf.table.rows)
        self.log.append(f"slide {n} tabela: termina em {(gf.top + gf.height) / EMU:.2f} in")

    # título curto + frase de abertura (design.md, slide 5 de 2026-10-19) --------
    def frase_abertura(self, n, velho, titulo, frase, destaque=None):
        """Troca o título do slide n (que começa por `velho`) por `titulo` curto e põe
        `frase` como 2º parágrafo: 20 pt 1E3A8A, `destaque` (trecho da frase) em negrito,
        barra 2563EB à esquerda. Implementação em montagem.py."""
        from montagem import frase_abertura
        frase_abertura(self.S[n - 1], velho, titulo, frase, destaque)
        self.log.append(f"slide {n}: título '{titulo}' + frase de abertura")

    # execução --------------------------------------------------------------
    def rodar(self, escopo):
        """gen_x.py [grupos,separados] [saída]: roda os g_* do script e grava."""
        todos = [k[2:] for k, v in escopo.items() if k.startswith("g_") and callable(v)]
        grupos = sys.argv[1].split(",") if len(sys.argv) > 1 else todos
        saida = sys.argv[2] if len(sys.argv) > 2 else f"{self.pasta}\\montado.pptx"
        for g in grupos:
            escopo["g_" + g](self)
        self.prs.save(saida)
        print("\n".join(self.log))
        print("OK", grupos, saida)
