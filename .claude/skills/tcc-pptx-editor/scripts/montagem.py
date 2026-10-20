"""Montagem de slide nos gen_*.py: texto com subscrito, caixas e cópia de formas.

Extraído dos gen_s10/gen_pu de 2026-10-19 (sintonia do PLL em etapas), que
reescreviam as mesmas funções. Edicao (edicoes.py) cuida do que é conferido
contra o deck; aqui fica o que posiciona e escreve.

    from montagem import caixa, copia, duplica_grupo, escreve, tira

    escreve(sh, ["1. Tensão em pu", "v_{abc} ÷ 16.329,93 V: V_{m} = 1 pu"])
    caixa(sh, 1.10, 2.20, 4.30, 0.81)            # polegadas
"""
import copy
import re

from pptx.util import Inches

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def runs_de(s):
    """'v_{q} ≈ ω_{n}^{2}' -> [(texto, None|'sub'|'sup'), ...]."""
    out, pos = [], 0
    for m in re.finditer(r"([_^])\{([^}]*)\}", s):
        if m.start() > pos:
            out.append((s[pos:m.start()], None))
        out.append((m.group(2), "sub" if m.group(1) == "_" else "sup"))
        pos = m.end()
    if pos < len(s):
        out.append((s[pos:], None))
    return out


def escreve(sh, paras):
    """Um parágrafo por item, no estilo do 1º run de cada parágrafo existente
    (o último é copiado se faltar parágrafo). Marcação _{x} = subscrito
    (baseline -25000), ^{x} = sobrescrito (30000), como no resto do deck."""
    ps = [p._p for p in sh.text_frame.paragraphs]
    while len(ps) < len(paras):
        n = copy.deepcopy(ps[-1])
        ps[-1].addnext(n)
        ps.append(n)
    for extra in ps[len(paras):]:
        extra.getparent().remove(extra)
    for p, texto in zip(ps, paras):
        rs = p.findall(A + "r")
        r0 = copy.deepcopy(rs[0])
        for r in rs:
            p.remove(r)
        end = p.find(A + "endParaRPr")
        for t, k in runs_de(texto):
            r = copy.deepcopy(r0)
            r.find(A + "t").text = t
            rpr = r.find(A + "rPr")
            rpr.attrib.pop("baseline", None)
            if k:
                rpr.set("baseline", "-25000" if k == "sub" else "30000")
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)


def caixa(sh, x, y, w, h):
    """Posição e tamanho em polegadas."""
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


def tira(sh):
    sh._element.getparent().remove(sh._element)


def copia(origem, destino):
    """Cópia de uma forma (de qualquer slide) no fim do slide destino. Imagem não:
    a mídia precisa de relate_to (ver Edicao.duplica_slide). Rodar renumera_ids."""
    el = copy.deepcopy(origem._element)
    destino.shapes._spTree.append(el)
    return destino.shapes[-1]


def duplica_grupo(formas_grupo):
    """Cópia de um conjunto de formas (ex.: caixa + título + corpo de uma etapa),
    inserida logo antes da 1ª, na mesma ordem. Os índices de slide.shapes de
    tudo que vem depois sobem len(formas_grupo). Rodar renumera_ids."""
    alvo = formas_grupo[0]._element
    for sh in formas_grupo:
        alvo.addprevious(copy.deepcopy(sh._element))


def frase_abertura(s, velho, titulo, frase, destaque=None):
    """Título curto + frase de abertura (design.md); ver Edicao.frase_abertura."""
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx_util import EMU
    from edicoes import Edicao
    c = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith(velho)]
    assert len(c) == 1, f"{len(c)} títulos começando por '{velho}'"
    t = c[0]
    ps = t.text_frame.paragraphs
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    p0 = ps[0]
    p0.runs[0].text = titulo
    for r in p0.runs[1:]:
        Edicao.tira_run(r)
    p1 = copy.deepcopy(p0._p)
    p0._p.addnext(p1)
    r0 = p1.findall(A + "r")[0]
    partes = [frase] if not destaque else frase.partition(destaque)
    assert not destaque or partes[1], f"'{destaque}' fora da frase"
    for i, txt in enumerate(partes):
        if not txt:
            continue
        r = copy.deepcopy(r0)
        r.find(A + "t").text = txt
        rpr = r.find(A + "rPr")
        rpr.set("sz", "2000")
        rpr.set("b", "1" if i == 1 else "0")
        for f in rpr.findall(A + "solidFill"):
            rpr.remove(f)
        sf = etree.Element(A + "solidFill")
        etree.SubElement(sf, A + "srgbClr").set("val", "1E3A8A")
        rpr.insert(0, sf)
        r0.addprevious(r)
    p1.remove(r0)
    end = p1.find(A + "endParaRPr")
    if end is not None:
        end.set("sz", "2000")
    ppr = p1.find(A + "pPr")
    ppr.set("marL", str(int(Inches(0.22))))
    sb = etree.Element(A + "spcBef")
    etree.SubElement(sb, A + "spcPts").set("val", "600")
    ppr.insert(0, sb)
    y = t.top / EMU + 0.58
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(t.left / EMU), Inches(y),
                           Inches(0.07), Inches(0.36))
    b.name = "Barra da frase"
    b.fill.solid()
    b.fill.fore_color.rgb = RGBColor.from_string("2563EB")
    b.line.fill.background()
    b.shadow.inherit = False
