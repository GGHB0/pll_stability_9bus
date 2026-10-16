"""Utilidades comuns aos scripts da skill tcc-pptx-editor (python-pptx).

Importar com:
    sys.path.insert(0, r"C:\\projetos\\pll_stability_9bus\\.claude\\skills\\tcc-pptx-editor\\scripts")
    from pptx_util import abrir, formas, texto_forma, notas, faixa

Unidades: tudo em polegadas (EMU / 914400) para bater com o design.md.
"""
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400


def utf8():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def abrir(path):
    return Presentation(path)


def pol(v):
    return (v or 0) / EMU


def faixa(spec, n):
    """'3-7,10' -> [3,4,5,6,7,10]; None -> 1..n (numeração do PowerPoint, base 1)."""
    if not spec:
        return list(range(1, n + 1))
    out = []
    for parte in spec.split(","):
        if "-" in parte:
            a, b = parte.split("-")
            out += range(int(a), int(b) + 1)
        else:
            out.append(int(parte))
    return [i for i in out if 1 <= i <= n]


def formas(container, prof=0):
    """Percorre as formas, entrando em grupos. Gera (forma, profundidade)."""
    for sh in container.shapes if hasattr(container, "shapes") else container:
        yield sh, prof
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from formas(sh.shapes, prof + 1)


def texto_forma(sh, sep=" | "):
    """Texto visível de uma forma: caixa de texto ou tabela (linhas por ' || ')."""
    if sh.has_text_frame:
        return sep.join(p.text for p in sh.text_frame.paragraphs)
    if getattr(sh, "has_table", False) and sh.has_table:
        linhas = []
        for row in sh.table.rows:
            linhas.append(" ; ".join(c.text.replace("\n", " ") for c in row.cells))
        return " || ".join(linhas)
    return ""


def tipo(sh):
    if getattr(sh, "has_table", False) and sh.has_table:
        return "tabela"
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "imagem"
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "grupo"
    if sh.has_text_frame and sh.text_frame.text.strip():
        return "texto"
    return "forma"


def notas(slide):
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text.strip()


_FALA = re.compile(r"\[(Bruno|Victor)\]\s*")


def falas(texto):
    """Divide as notas em [(orador, trecho)]. Texto antes do 1º marcador -> ('?', ...).
    Remove o número de slide solto que alguns geradores deixam no fim."""
    texto = re.sub(r"\s+\d{1,2}\s*$", "", texto)
    partes = _FALA.split(texto)
    out = []
    if partes[0].strip():
        out.append(("?", partes[0].strip()))
    for i in range(1, len(partes) - 1, 2):
        out.append((partes[i], partes[i + 1].strip()))
    return out
