"""Peças reutilizáveis para trocar o visual do deck (casos/troca_template.md).

Extraído do gen_template.py de 2026-10-16 (template "Project Portfolio"). Uso:

    from template_util import Rascunho, desloca, rotulo_secao, estilo_chapado
    r = Rascunho(prs)                     # slide temporário para fabricar formas
    tree.append(copy.deepcopy(r.caixa("UERJ", 0.85, 0.12, 4.2, 0.4, 12)._element))
    ...
    r.descarta()                          # SEMPRE antes do prs.save

desloca() NÃO é idempotente: roda uma vez sobre o deck do staging, nunca sobre a própria saída.
"""
import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from pptx_util import EMU

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
TEXTO = RGBColor(0x0F, 0x17, 0x2A)


def i(v):
    return int(v * EMU)


class Rascunho:
    """LayoutShapes não tem add_textbox/add_connector: fabrica num slide temporário e copia o XML."""

    def __init__(self, prs):
        self.prs = prs
        self.s = prs.slides.add_slide(prs.slide_layouts[0])

    def caixa(self, texto, x, y, w, h, pt, bold=False, cor=TEXTO, alinha=PP_ALIGN.LEFT, fonte="Arial", rot=0):
        tb = self.s.shapes.add_textbox(i(x), i(y), i(w), i(h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = alinha
        r = p.add_run()
        r.text = texto
        r.font.size, r.font.bold, r.font.name = Pt(pt), bold, fonte
        r.font.color.rgb = cor
        tb.rotation = rot
        return tb

    def linha(self, x1, y1, x2, y2, pt=1.0, cor=TEXTO):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, i(x1), i(y1), i(x2), i(y2))
        c.line.color.rgb = cor
        c.line.width = Pt(pt)
        return c

    def descarta(self):
        lst = self.prs.slides._sldIdLst
        alvo = [x for x in lst if self.prs.part.related_part(x.rId) is self.s.part][0]
        rid = alvo.rId
        lst.remove(alvo)
        self.prs.part.drop_rel(rid)


def desloca(slides, x0_velho, larg_velha, x0_novo, larg_nova):
    """Leva o conteúdo de [x0_velho, +larg_velha] para [x0_novo, +larg_nova] em x.
    Imagem escala por igual (centrada no lugar); tabela escala as colunas (gridCol);
    número do slide fica. Retorna o fator K."""
    K = larg_nova / larg_velha
    for s in slides:
        for sh in s.shapes:
            if sh.name.startswith("Slide Number"):
                continue
            x, w = sh.left / EMU, sh.width / EMU
            if sh.shape_type == 13:
                h = sh.height / EMU
                nh = h * K
                sh.left, sh.width = i(x0_novo + (x - x0_velho) * K), i(w * K)
                sh.top, sh.height = i(sh.top / EMU + (h - nh) / 2), i(nh)
                continue
            sh.left, sh.width = i(x0_novo + (x - x0_velho) * K), i(w * K)
            if getattr(sh, "has_table", False) and sh.has_table:
                for gc in sh.table._tbl.tblGrid.gridCol_lst:
                    gc.w = int(gc.w * K)
    return K


def rotulo_secao(rasc, slide, nome, cx=0.40, cy=4.10, comp=5.8, alt=0.44, pt=18):
    """Nome da seção girado 270° centrado em (cx, cy); forma "Seção" no slide."""
    tb = rasc.caixa(nome, cx - comp / 2, cy - alt / 2, comp, alt, pt, bold=True, alinha=PP_ALIGN.CENTER, rot=270)
    slide.shapes._spTree.append(copy.deepcopy(tb._element))
    slide.shapes[-1].name = "Seção"


def cor(el):
    """('s', 'accent1') para schemeClr, ('r', 'FFFFFF') para srgbClr, None sem cor."""
    if el is None:
        return None
    c = el.find(A + "schemeClr")
    if c is not None:
        return ("s", c.get("val"))
    c = el.find(A + "srgbClr")
    return ("r", c.get("val").upper()) if c is not None else None


def pinta(solid, hexa):
    for c in list(solid):
        solid.remove(c)
    etree.SubElement(solid, A + "srgbClr").set("val", hexa)


def sem_borda(sppr):
    ln = sppr.find(A + "ln")
    if ln is None:
        ln = etree.SubElement(sppr, A + "ln")
    for c in list(ln):
        ln.remove(c)
    etree.SubElement(ln, A + "noFill")


def estilo_chapado(slides, troca_fill, tira_borda_de):
    """Canto reto em todo roundRect; troca_fill = {('s','accent6'): 'DBEAFE', ...} pinta o
    preenchimento (de forma e de célula) e tira a borda dessas formas; tira_borda_de = conjunto de
    cores de borda removidas de formas preenchidas (ex.: {('s','accent4')}). Retorna contagens."""
    n = {"rr": 0, "ln": 0, "fill": 0}
    for s in slides:
        for el in s._element.iter(A + "prstGeom"):
            if el.get("prst") == "roundRect":
                el.set("prst", "rect")
                n["rr"] += 1
        for sppr in s._element.iter(P + "spPr"):
            solid = sppr.find(A + "solidFill")
            f = cor(solid)
            lc = cor(sppr.find(A + "ln/" + A + "solidFill"))
            if f in troca_fill:
                pinta(solid, troca_fill[f])
                n["fill"] += 1
                if lc is not None:
                    sem_borda(sppr)
                    n["ln"] += 1
            elif f is not None and lc in tira_borda_de:
                sem_borda(sppr)
                n["ln"] += 1
        for tcpr in s._element.iter(A + "tcPr"):
            solid = tcpr.find(A + "solidFill")
            if cor(solid) in troca_fill:
                pinta(solid, troca_fill[cor(solid)])
                n["fill"] += 1
    return n
